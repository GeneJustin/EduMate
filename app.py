import os
import json
import time
from datetime import datetime
from flask import Flask, render_template, request, jsonify
from dotenv import load_dotenv
from google import genai
from supabase import create_client
from pypdf import PdfReader
from pptx import Presentation
from werkzeug.utils import secure_filename

load_dotenv()

app = Flask(__name__)

UPLOAD_FOLDER = "/tmp/uploads"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER


supabase_url = os.getenv("SUPABASE_URL")
supabase_key = os.getenv("SUPABASE_KEY")

if not supabase_url or not supabase_key:
    raise ValueError("SUPABASE_URL and SUPABASE_KEY are required.")

supabase = create_client(
    supabase_url,
    supabase_key
)


api_keys = [
    os.getenv("GEMINI_API_KEY"),
    os.getenv("GEMINI_API_KEY2"),
    os.getenv("GEMINI_API_KEY3"),
    os.getenv("GEMINI_API_KEY4")
]

api_keys = [
    key for key in api_keys
    if key
]

if not api_keys:
    raise ValueError("No Gemini API keys found.")

clients = [
    genai.Client(api_key=key)
    for key in api_keys
]


def generate_json(prompt, schema):

    last_error = None

    for client_index, current_client in enumerate(clients):

        for attempt in range(2):

            try:

                response = current_client.models.generate_content(
                    model="gemini-3.6-flash",
                    contents=prompt,
                    config={
                        "response_mime_type": "application/json",
                        "response_schema": schema
                    }
                )

                if not response.text:
                    raise ValueError(
                        "Gemini returned an empty response."
                    )

                return json.loads(response.text)

            except Exception as e:

                last_error = e
                error_message = str(e)

                if (
                    "429" in error_message
                    or "RESOURCE_EXHAUSTED" in error_message
                ):

                    print(
                        f"Gemini key {client_index + 1} quota exceeded."
                    )

                    break

                if (
                    "503" in error_message
                    or "UNAVAILABLE" in error_message
                ):

                    if attempt == 0:

                        print(
                            f"Gemini key {client_index + 1} unavailable. Retrying..."
                        )

                        time.sleep(2)

                        continue

                    break

                raise

    raise last_error


def generate_text(prompt):

    last_error = None

    for client_index, current_client in enumerate(clients):

        for attempt in range(2):

            try:

                response = current_client.models.generate_content(
                    model="gemini-3.6-flash",
                    contents=prompt
                )

                if not response.text:
                    raise ValueError(
                        "Gemini returned an empty response."
                    )

                return response.text

            except Exception as e:

                last_error = e
                error_message = str(e)

                if (
                    "429" in error_message
                    or "RESOURCE_EXHAUSTED" in error_message
                ):

                    print(
                        f"Gemini key {client_index + 1} quota exceeded."
                    )

                    break

                if (
                    "503" in error_message
                    or "UNAVAILABLE" in error_message
                ):

                    if attempt == 0:

                        print(
                            f"Gemini key {client_index + 1} unavailable. Retrying..."
                        )

                        time.sleep(2)

                        continue

                    break

                raise

    raise last_error


def extract_pdf(path):

    reader = PdfReader(path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def extract_ppt(path):

    presentation = Presentation(path)

    text = ""

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        text += f"\n--- Slide {slide_number} ---\n"

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text += shape.text + "\n"

    return text


def extract_text(path, file_type):

    if file_type == "pdf":
        return extract_pdf(path)

    if file_type in ["ppt", "pptx"]:
        return extract_ppt(path)

    return ""


def generate_summary(content):

    prompt = f"""
You are an AI study assistant.

Analyze the learning material below and create a concise study summary.

IMPORTANT:
- Return ONLY valid JSON.
- Do not include any introduction.
- Do not include any conclusion outside the JSON.
- Do not write phrases such as "Here is the summary".
- Focus only on important information.
- Use clear Indonesian.
- Keep explanations concise but useful for university students.
- Do not add information that is not present in the material.
- If there are no formulas, return an empty array for formulas.
- If there is no exam-specific information, infer exam focus only from important concepts explicitly present in the material.

The JSON structure must contain:
- title
- overview
- key_points
- important_concepts
- formulas
- exam_focus

Learning material:

{content}
"""

    schema = {
        "type": "object",
        "properties": {
            "title": {
                "type": "string"
            },
            "overview": {
                "type": "string"
            },
            "key_points": {
                "type": "array",
                "items": {
                    "type": "string"
                }
            },
            "important_concepts": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "concept": {
                            "type": "string"
                        },
                        "explanation": {
                            "type": "string"
                        }
                    },
                    "required": [
                        "concept",
                        "explanation"
                    ]
                }
            },
            "formulas": {
                "type": "array",
                "items": {
                    "type": "string"
                }
            },
            "exam_focus": {
                "type": "array",
                "items": {
                    "type": "string"
                }
            }
        },
        "required": [
            "title",
            "overview",
            "key_points",
            "important_concepts",
            "formulas",
            "exam_focus"
        ]
    }

    return generate_json(prompt, schema)


def parse_summary(summary, title):

    if not summary:
        return {
            "title": title,
            "overview": "",
            "key_points": [],
            "important_concepts": [],
            "formulas": [],
            "exam_focus": []
        }

    try:

        return json.loads(summary)

    except (json.JSONDecodeError, TypeError):

        return {
            "title": title,
            "overview": summary,
            "key_points": [],
            "important_concepts": [],
            "formulas": [],
            "exam_focus": []
        }


@app.route("/")
def index():

    return render_template("index.html")


@app.route("/chat")
def chat():

    return render_template("chat.html")


@app.route("/materials")
def materials():

    try:

        response = (
            supabase
            .table("materials")
            .select("*")
            .order("id", desc=True)
            .execute()
        )

        materials_data = response.data

        return render_template(
            "materials.html",
            materials=materials_data
        )

    except Exception as e:

        print("MATERIALS PAGE ERROR:", repr(e))

        return render_template(
            "materials.html",
            materials=[]
        )


@app.route("/quiz")
def quiz():

    try:

        response = (
            supabase
            .table("materials")
            .select("*")
            .order("id", desc=True)
            .execute()
        )

        materials_data = response.data

        return render_template(
            "quiz.html",
            materials=materials_data
        )

    except Exception as e:

        print("QUIZ PAGE ERROR:", repr(e))

        return render_template(
            "quiz.html",
            materials=[]
        )


@app.route("/tracker")
def tracker():

    return render_template("tracker.html")


@app.route("/api/chat", methods=["POST"])
def api_chat():

    try:

        data = request.get_json()

        if not data:
            return jsonify({
                "error": "Invalid request."
            }), 400

        message = data.get("message", "").strip()

        if not message:
            return jsonify({
                "error": "Message is empty."
            }), 400

        prompt = f"""
You are StudyMate, an AI learning assistant.

Your purpose is to help students understand academic topics.

Rules:
- Answer in clear Indonesian.
- Explain concepts step by step.
- Give examples when useful.
- If the student asks for code, explain the code.
- Do not simply give an answer when explanation would help learning.
- Keep answers reasonably concise.

Student question:

{message}
"""

        response = generate_text(prompt)

        return jsonify({
            "response": response
        })

    except Exception as e:

        print("CHAT ERROR:", repr(e))

        return jsonify({
            "error": str(e)
        }), 500


@app.route("/api/upload", methods=["POST"])
def upload_material():

    try:

        if "file" not in request.files:
            return jsonify({
                "error": "No file uploaded."
            }), 400

        file = request.files["file"]

        if file.filename == "":
            return jsonify({
                "error": "No file selected."
            }), 400

        extension = (
            file.filename
            .rsplit(".", 1)[-1]
            .lower()
        )

        if extension not in ["pdf", "ppt", "pptx"]:
            return jsonify({
                "error": "Only PDF, PPT, and PPTX are supported."
            }), 400

        filename = secure_filename(file.filename)

        if not filename:
            return jsonify({
                "error": "Invalid filename."
            }), 400

        path = os.path.join(
            app.config["UPLOAD_FOLDER"],
            filename
        )

        file.save(path)

        print("FILE SAVED:", path)

        content = extract_text(
            path,
            extension
        )

        print(
            "TEXT EXTRACTED:",
            len(content),
            "characters"
        )

        if not content.strip():
            return jsonify({
                "error": "Could not extract text from this file."
            }), 400

        summary = generate_summary(content)

        print("SUMMARY GENERATED")

        summary_json = json.dumps(
            summary,
            ensure_ascii=False
        )

        response = (
            supabase
            .table("materials")
            .insert({
                "filename": filename,
                "title": os.path.splitext(filename)[0],
                "file_type": extension,
                "summary": summary_json,
                "content": content,
                "created_at": datetime.now().isoformat()
            })
            .execute()
        )

        print("MATERIAL SAVED TO SUPABASE")

        return jsonify({
            "success": True,
            "filename": filename,
            "summary": summary
        })

    except Exception as e:

        print("UPLOAD ERROR:", repr(e))

        return jsonify({
            "error": str(e)
        }), 500


@app.route("/api/materials/<int:material_id>")
def get_material(material_id):

    try:

        response = (
            supabase
            .table("materials")
            .select("*")
            .eq("id", material_id)
            .single()
            .execute()
        )

        material = response.data

        if not material:
            return jsonify({
                "error": "Material not found."
            }), 404

        summary = parse_summary(
            material.get("summary"),
            material.get("title")
        )

        return jsonify({
            "id": material.get("id"),
            "title": material.get("title"),
            "filename": material.get("filename"),
            "file_type": material.get("file_type"),
            "summary": summary
        })

    except Exception as e:

        print("MATERIAL ERROR:", repr(e))

        return jsonify({
            "error": str(e)
        }), 500


@app.route("/api/quiz", methods=["POST"])
def generate_quiz():

    try:

        data = request.get_json()

        if not data:
            return jsonify({
                "error": "Invalid request."
            }), 400

        material_id = data.get("material_id")

        try:
            num_questions = int(
                data.get("num_questions", 5)
            )
        except (ValueError, TypeError):
            return jsonify({
                "error": "Invalid number of questions."
            }), 400

        if not material_id:
            return jsonify({
                "error": "Please select a material."
            }), 400

        if num_questions < 1 or num_questions > 20:
            return jsonify({
                "error": "Number of questions must be between 1 and 20."
            }), 400

        response = (
            supabase
            .table("materials")
            .select("title, content")
            .eq("id", material_id)
            .single()
            .execute()
        )

        material = response.data

        if not material:
            return jsonify({
                "error": "Material not found."
            }), 404

        title = material["title"]
        content = material["content"]

        prompt = f"""
You are an AI quiz generator for university students.

Create exactly {num_questions} multiple-choice questions
based ONLY on the learning material below.

IMPORTANT:
- Return ONLY valid JSON.
- Do not write any introduction.
- Do not write any conclusion.
- Do not write phrases such as "Here are the questions".
- Do not include markdown.
- Do not add information outside the material.
- Every question must have exactly four options.
- There must be exactly one correct answer.
- The explanation must briefly explain why the answer is correct.
- Questions should test understanding, not only memorization.
- Use clear Indonesian.
- Make the options plausible.
- Avoid duplicate questions.
- Return exactly {num_questions} questions.

Learning material title:
{title}

Learning material:
{content}
"""

        schema = {
            "type": "object",
            "properties": {
                "title": {
                    "type": "string"
                },
                "questions": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "question": {
                                "type": "string"
                            },
                            "options": {
                                "type": "object",
                                "properties": {
                                    "A": {
                                        "type": "string"
                                    },
                                    "B": {
                                        "type": "string"
                                    },
                                    "C": {
                                        "type": "string"
                                    },
                                    "D": {
                                        "type": "string"
                                    }
                                },
                                "required": [
                                    "A",
                                    "B",
                                    "C",
                                    "D"
                                ]
                            },
                            "answer": {
                                "type": "string",
                                "enum": [
                                    "A",
                                    "B",
                                    "C",
                                    "D"
                                ]
                            },
                            "explanation": {
                                "type": "string"
                            }
                        },
                        "required": [
                            "question",
                            "options",
                            "answer",
                            "explanation"
                        ]
                    }
                }
            },
            "required": [
                "title",
                "questions"
            ]
        }

        quiz = generate_json(
            prompt,
            schema
        )

        if len(quiz.get("questions", [])) != num_questions:

            return jsonify({
                "error": "Gemini did not generate the requested number of questions."
            }), 500

        return jsonify(quiz)

    except Exception as e:

        print("QUIZ ERROR:", repr(e))

        return jsonify({
            "error": str(e)
        }), 500


@app.route("/api/study/start", methods=["POST"])
def start_study():

    try:

        start_time = datetime.now().isoformat()

        response = (
            supabase
            .table("study_sessions")
            .insert({
                "start_time": start_time,
                "created_at": start_time
            })
            .select("id, start_time")
            .execute()
        )

        session = response.data[0]

        return jsonify({
            "session_id": session["id"],
            "start_time": session["start_time"]
        })

    except Exception as e:

        print("STUDY START ERROR:", repr(e))

        return jsonify({
            "error": str(e)
        }), 500


@app.route("/api/study/stop", methods=["POST"])
def stop_study():

    try:

        data = request.get_json()

        if not data:
            return jsonify({
                "error": "Invalid request."
            }), 400

        session_id = data.get("session_id")

        if not session_id:
            return jsonify({
                "error": "Session ID is required."
            }), 400

        response = (
            supabase
            .table("study_sessions")
            .select("*")
            .eq("id", session_id)
            .single()
            .execute()
        )

        session = response.data

        if not session:

            return jsonify({
                "error": "Session not found."
            }), 404

        start = datetime.fromisoformat(
            session["start_time"].replace("Z", "+00:00")
        )

        end = datetime.now(start.tzinfo)

        duration = int(
            (end - start).total_seconds()
        )

        (
            supabase
            .table("study_sessions")
            .update({
                "end_time": end.isoformat(),
                "duration": duration
            })
            .eq("id", session_id)
            .execute()
        )

        return jsonify({
            "duration": duration
        })

    except Exception as e:

        print("STUDY STOP ERROR:", repr(e))

        return jsonify({
            "error": str(e)
        }), 500


@app.route("/api/study/stats")
def study_stats():

    try:

        response = (
            supabase
            .table("study_sessions")
            .select("duration")
            .gt("duration", 0)
            .execute()
        )

        sessions = response.data

        total_seconds = sum(
            session.get("duration", 0) or 0
            for session in sessions
        )

        return jsonify({
            "sessions": len(sessions),
            "total_seconds": total_seconds
        })

    except Exception as e:

        print("STUDY STATS ERROR:", repr(e))

        return jsonify({
            "error": str(e)
        }), 500


if __name__ == "__main__":

    app.run(
        debug=True
    )
